"""CGJR standardized tables: employment table (G16) and historical macro table (G14).

Both builders return a :class:`~cgjrkit.core.provenance.Result` whose
``data`` is the *display* table (one row per table row, one column per
table column) rather than a canonical long frame; ``meta["table"]``
describes the layout (sections, column groups) used by :func:`to_xlsx` and
:func:`to_html`, and ``meta["cells"]`` lists every cell with its column
entity, period and, for peer cells, the number of members with data
(JSON-safe, for binding numbers in drafts). Provenance and caveats of all
inputs are carried forward as in every other method.

Standardized employment table (Morocco Annex 2, North Macedonia p. 96,
Philippines p. 98, Indonesia Table A.2)
    Section "growth": annualized employment growth by sector for the focus
    country over several windows (e.g. 1991-2009, 2010-2022, 2020, 2021,
    2022) and, for each peer set, the peer statistic (median by default;
    Morocco and North Macedonia print averages: ``stat="mean"``) over the
    peer windows. Section "share": sector shares of total employment in one
    year, focus vs peer statistic.

Standardized macro table (historical columns)
    One row per indicator (GDP growth, GDP per capita growth, investment
    % of GDP, ...), period averages for the focus country and peer
    statistics of the members' period averages.

Data column keys
    Employment table: ``"<section>|<entity>|<period>"`` (``"growth|MAR|2010-2022"``,
    ``"share|STRUCT|2022"``). Macro table: ``"<entity>|<period>"``. The
    entity is the focus ISO3 or the peer-set code. Label columns are
    ``section`` and ``row`` (plus ``source`` in the macro table).

Window conventions
    Macro table: a window ``(start, end)`` is the inclusive set of years
    ``start..end`` for *every* aggregation. ``agg="mean"`` / ``"sum"``
    use the observations dated ``start..end``; ``agg="cagr"`` compounds the
    growth of the same years, i.e. between the levels of ``start - 1`` and
    ``end`` (as :func:`cgjrkit.methods.harmonize.period_average` notes, a CAGR
    window starts one year earlier than the matching average of annual growth
    rates). A single period ``"2020"`` is the window ``("2020", "2020")``:
    growth from 2019 to 2020. So a mean-of-growth row and a CAGR row in the
    same column cover the same growth years.

    Employment table: every cell is a growth rate between *levels*, following
    the standardized employment tables, whose employment series (ILO modelled
    estimates) start in 1991: a window ``(start, end)`` is the CAGR between
    the levels at ``start`` and ``end`` (growth years ``start+1..end``) and a
    single period ``"2020"`` is growth from 2019 to 2020. Consecutive windows
    such as ``"1991-2009"`` and ``"2010-2022"`` therefore leave the 2009->2010
    change out of both columns; a ``window_change_not_covered`` info caveat
    lists such changes.
"""

from __future__ import annotations

import html
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal, Mapping, Sequence

import numpy as np
import pandas as pd

from cgjrkit.core.caveats import Caveat
from cgjrkit.core.provenance import Result
from cgjrkit.methods.harmonize import (
    FrequencyMismatchError,
    PeriodSpec,
    _limited,
    _window,
    combine_status,
    format_period,
    parse_period,
)
from cgjrkit.methods.jobs_shapley import _endpoints
from cgjrkit.methods.peers import PeerSet, peer_aggregate
from cgjrkit.methods.structural_change import (
    SectorMap,
    StructuralChangeError,
    _done,
    _endpoint_caveats,
    _modelled_caveat,
    _single_series,
    _years_between,
    sector_panel,
)

__all__ = [
    "MacroRow",
    "standardized_employment_table",
    "standardized_macro_table",
    "to_xlsx",
    "to_html",
]

Frame = Result | pd.DataFrame
PeerSets = Mapping[str, "PeerSet | Sequence[str]"] | Sequence[PeerSet]
Stat = Literal["median", "mean"]

_HEADER_FILL = "D9E2EA"
_SECTION_FILL = "EEF1F4"
_MISSING = ".."


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def _peersets(peer_sets: PeerSets) -> list[PeerSet]:
    if isinstance(peer_sets, PeerSet):
        out = [peer_sets]
    elif isinstance(peer_sets, Mapping):
        out = [spec if isinstance(spec, PeerSet) else PeerSet(name=str(label), members=list(spec))
               for label, spec in peer_sets.items()]
    else:
        out = list(peer_sets)
        if not all(isinstance(p, PeerSet) for p in out):
            raise TypeError("peer_sets must be PeerSet objects or a mapping {name: PeerSet | [ISO3, ...]}")
    codes = [p.code for p in out]
    if len(set(codes)) != len(codes):
        raise ValueError(f"peer-set codes must be unique, got {codes}")
    return out


def _growth(x0: float, xT: float, years: float, kind: str) -> float:
    if not (np.isfinite(x0) and np.isfinite(xT)) or x0 <= 0 or xT <= 0:
        return float("nan")
    if kind == "cagr":
        return 100.0 * ((xT / x0) ** (1.0 / years) - 1.0)
    return 100.0 * math.log(xT / x0) / years


_FY_HINT = ("convert first with harmonize.fiscal_to_calendar() or harmonize.calendar_to_fiscal(), or give windows "
            "in the data's frequency")


def _period_freqs(periods: Sequence[str]) -> set[str]:
    out = set()
    for p in pd.unique(pd.Series(list(periods), dtype=object)):
        try:
            out.add(parse_period(str(p)).freq)
        except ValueError:  # e.g. a period-average label such as "2011-2019"
            continue
    return out


def _check_frequency(periods: Sequence[str], window_freqs: Mapping[str, str], what: str) -> None:
    """Data periods must have one frequency, the same as every window's (FY2011/12 is not 2011)."""
    freqs = _period_freqs(periods)
    if len(freqs) > 1:
        raise FrequencyMismatchError(f"{what} mixes period frequencies {sorted(freqs)}; {_FY_HINT}")
    if not freqs:
        return
    f = next(iter(freqs))
    bad = {lab: wf for lab, wf in window_freqs.items() if wf != f}
    if bad:
        raise FrequencyMismatchError(f"{what} has {f} periods but windows {sorted(bad)} are {sorted(set(bad.values()))}; "
                                     f"{_FY_HINT}")


def _uncovered_changes(windows: Mapping[str, tuple[str, str]]) -> list[str]:
    """Level-to-level changes inside the span of the windows that no window covers (``"2009->2010"``)."""
    out = []
    by_freq: dict[str, set[int]] = {}
    for lo, hi in windows.values():
        a, b = parse_period(lo), parse_period(hi)
        by_freq.setdefault(a.freq, set()).update(range(a.ordinal + 1, b.ordinal + 1))
    for freq, covered in sorted(by_freq.items()):
        for o in range(min(covered), max(covered) + 1):
            if o not in covered:
                out.append(f"{format_period(freq, o - 1)}->{format_period(freq, o)}")
    return out


def _check_stat(stat: str) -> None:
    if stat not in ("median", "mean"):
        raise ValueError("stat must be 'median' or 'mean'")


def _peer_values(rows: list[dict], ps: PeerSet, stat: str, min_n: int) -> tuple[dict, dict, Result | None]:
    """Run :func:`peer_aggregate` on member values; return ``({(indicator, period): value}, {..: n}, result)``."""
    if not rows:
        return {}, {}, None
    frame = pd.DataFrame(rows)
    frame = frame[frame["iso3"].isin(ps.members)]
    if frame.empty:
        return {}, {}, None
    agg = peer_aggregate(frame, ps, stats=[stat], min_n=min_n)
    vals = {(r.indicator, r.period): float(r.value) for r in agg.data.itertuples()}
    ns = {(r.indicator, r.period): int(r.n) for r in agg.data.itertuples()}
    return vals, ns, agg


def _member_row(iso3: str, indicator: str, period: str, value: float, unit: str, status: str = "") -> dict:
    return {"source": "cgjrkit", "dataset": "standard_tables", "indicator": indicator, "indicator_label": indicator,
            "iso3": iso3, "country": "", "period": period, "value": value, "unit": unit, "obs_status": status}


def _country_name(df: pd.DataFrame, iso3: str) -> str:
    names = [c for c in df.loc[df["iso3"] == iso3, "country"] if c]
    return names[0] if names else iso3


def _assemble(table_rows: list[dict], columns: list[dict], label_cols: list[str]) -> pd.DataFrame:
    keys = [c["key"] for c in columns]
    df = pd.DataFrame(table_rows)
    for k in keys:
        if k not in df.columns:
            df[k] = np.nan
    df = df[label_cols + keys]
    for k in keys:
        df[k] = df[k].astype("float64")
    return df.reset_index(drop=True)


def _cells(data: pd.DataFrame, columns: list[dict], ns: Mapping[tuple[str, str], int]) -> list[dict]:
    out = []
    for _, r in data.iterrows():
        for c in columns:
            v = r[c["key"]]
            if pd.isna(v):
                continue
            cell = {"section": r["section"], "row": r["row"], "column": c["key"], "entity": c["entity"],
                    "period": c["sub"], "value": float(v)}
            n = ns.get((c["key"], r["row"]))
            if n is not None:
                cell["n"] = n
            out.append(cell)
    return out


# ---------------------------------------------------------------------------
# employment table
# ---------------------------------------------------------------------------
def standardized_employment_table(
    emp_by_sector: Frame,
    focus: str,
    peer_sets: PeerSets,
    periods: Mapping[str, PeriodSpec],
    *,
    sectors: SectorMap | None = None,
    share_year: str | None = None,
    peer_periods: Sequence[str] | None = None,
    total: str | None = None,
    total_label: str = "Total",
    stat: Stat = "median",
    growth: Literal["cagr", "log"] = "cagr",
    min_n: int = 3,
    title: str | None = None,
    source_note: str = "",
) -> Result:
    """Standardized employment table: employment growth and sector shares, focus vs peers.

    ``emp_by_sector`` is employment by sector (ILO modelled estimates, LFS)
    for the focus and peer countries; ``sectors`` maps indicator codes to row
    labels (in row order). ``total`` names a sector that is total
    employment; otherwise the total is the sum of the sectors (missing when
    any sector is missing).

    ``periods``: ``{label: (start, end) | single period}`` for the focus
    growth columns, e.g. ``{"1991-2009": ("1991", "2009"), "2010-2022":
    ("2010", "2022"), "2020": "2020"}``; ``peer_periods`` (labels, default
    all) are the windows shown for each peer set. Growth is the CAGR between
    the window's end-point levels (``growth="log"`` for log growth): growth
    years ``start+1..end``, so ``"1991-2009"`` and ``"2010-2022"`` leave
    the 2009->2010 change out of both (``window_change_not_covered`` info
    caveat; use ``("2009", "2022")`` to include it). Data and windows must
    share one frequency (:class:`~cgjrkit.methods.harmonize.FrequencyMismatchError`).
    ``share_year`` (default: end of the last window) is the year of the
    sector-share columns. Peer cells are the ``stat`` (median or mean) across
    members with data (:func:`cgjrkit.methods.peers.peer_aggregate`, which adds
    ``peer_small_n`` / ``peer_set_not_frozen`` caveats).
    """
    _check_stat(stat)
    if growth not in ("cagr", "log"):
        raise ValueError("growth must be 'cagr' or 'log'")
    if not periods:
        raise ValueError("periods must name at least one window")
    focus = str(focus).strip().upper()
    psets = _peersets(peer_sets)
    windows = {str(lab): _endpoints(spec) for lab, spec in periods.items()}
    for lo, hi in windows.values():
        _years_between(lo, hi)
    plabels = list(windows) if peer_periods is None else [str(p) for p in peer_periods]
    unknown = [p for p in plabels if p not in windows]
    if unknown:
        raise ValueError(f"peer_periods {unknown} are not labels of periods {list(windows)}")
    share_year = str(share_year).strip() if share_year else list(windows.values())[-1][1]
    parse_period(share_year)

    res, panel, order = sector_panel(emp_by_sector, sectors, name="emp_by_sector")
    if total is not None and total not in order:
        raise StructuralChangeError(f"total={total!r} is not one of the sectors {order}")
    rows_sectors = [s for s in order if s != total]
    if focus not in set(panel["iso3"]):
        raise StructuralChangeError(f"no employment data for the focus country {focus}")
    relevant = {focus} | {m for ps in psets for m in ps.members}
    window_freqs = {lab: parse_period(lo).freq for lab, (lo, _) in windows.items()}
    window_freqs["share_year"] = parse_period(share_year).freq
    _check_frequency(panel.loc[panel["iso3"].isin(relevant), "period"].tolist(), window_freqs, "emp_by_sector")

    wide = panel.pivot_table(index=["iso3", "period"], columns="sector", values="value", aggfunc="first")
    status = panel.groupby(["iso3", "period"])["obs_status"].agg(combine_status)
    wide = wide.reindex(columns=order)
    if total is not None:
        tot = wide[total]
    else:
        tot = wide[rows_sectors].sum(axis=1, min_count=len(rows_sectors)).where(wide[rows_sectors].notna().all(axis=1))
    levels = wide[rows_sectors].copy()
    levels[total_label] = tot

    def level(iso3: str, period: str, sector: str) -> float:
        try:
            return float(levels.loc[(iso3, period), sector])
        except KeyError:
            return float("nan")

    def st(iso3: str, period: str) -> str:
        return str(status.get((iso3, period), ""))

    countries = sorted(set(panel["iso3"]))
    row_names = rows_sectors + [total_label]
    # growth and shares per country (long)
    growth_vals: dict[tuple[str, str, str], float] = {}
    share_vals: dict[tuple[str, str], float] = {}
    statuses: list[str] = []
    missing: list[str] = []
    for c in countries:
        for label, (lo, hi) in windows.items():
            yrs = _years_between(lo, hi)
            for s in row_names:
                g = _growth(level(c, lo, s), level(c, hi, s), yrs, growth)
                growth_vals[(c, s, label)] = g
                if c == focus and not np.isfinite(g):
                    missing.append(f"growth {s} {label}")
            statuses += [st(c, lo), st(c, hi)] if c == focus or any(c in p.members for p in psets) else []
        t = level(c, share_year, total_label)
        for s in rows_sectors:
            v = level(c, share_year, s)
            share_vals[(c, s)] = 100.0 * v / t if np.isfinite(v) and np.isfinite(t) and t > 0 else float("nan")
            if c == focus and not np.isfinite(share_vals[(c, s)]):
                missing.append(f"share {s} {share_year}")

    focus_name = _country_name(panel, focus)
    columns: list[dict] = []
    for label in windows:
        columns.append({"key": f"growth|{focus}|{label}", "group": focus_name, "sub": label, "sections": ["growth"],
                        "entity": focus, "kind": "focus"})
    for ps in psets:
        for label in plabels:
            columns.append({"key": f"growth|{ps.code}|{label}", "group": f"{ps.name} ({stat})", "sub": label,
                            "sections": ["growth"], "entity": ps.code, "kind": "peer"})
    columns.append({"key": f"share|{focus}|{share_year}", "group": focus_name, "sub": share_year,
                    "sections": ["share"], "entity": focus, "kind": "focus"})
    for ps in psets:
        columns.append({"key": f"share|{ps.code}|{share_year}", "group": f"{ps.name} ({stat})", "sub": share_year,
                        "sections": ["share"], "entity": ps.code, "kind": "peer"})

    # peer statistics
    peer_results: list[Result] = []
    peer_cells: dict[tuple[str, str], float] = {}
    peer_n: dict[tuple[str, str], int] = {}
    focus_in_peers = [ps.code for ps in psets if focus in ps.members]
    for ps in psets:
        rows = []
        for m in ps.members:
            for s in row_names:
                for label in plabels:
                    g = growth_vals.get((m, s, label), float("nan"))
                    if np.isfinite(g):
                        rows.append(_member_row(m, f"growth|{s}", label, g, "% per year"))
            for s in rows_sectors:
                v = share_vals.get((m, s), float("nan"))
                if np.isfinite(v):
                    rows.append(_member_row(m, f"share|{s}", share_year, v, "% of total employment"))
        vals, ns, agg = _peer_values(rows, ps, stat, min_n)
        if agg is None:
            agg = Result(data=pd.DataFrame(), caveats=[Caveat(
                "peer_no_data", "warn", f"No member of peer set {ps.name!r} has employment data for these windows.",
                {"peer_set": ps.code, "members": ps.members})])
        peer_results.append(agg)
        for (ind, per), v in vals.items():
            sec, s = ind.split("|", 1)
            key = f"{sec}|{ps.code}|{per}"
            peer_cells[(key, s)] = v
            peer_n[(key, s)] = ns[(ind, per)]

    table_rows: list[dict] = []
    for s in row_names:
        r: dict[str, Any] = {"section": "growth", "row": s}
        for label in windows:
            r[f"growth|{focus}|{label}"] = growth_vals[(focus, s, label)]
        for ps in psets:
            for label in plabels:
                key = f"growth|{ps.code}|{label}"
                r[key] = peer_cells.get((key, s), float("nan"))
        table_rows.append(r)
    for s in rows_sectors:
        r = {"section": "share", "row": s, f"share|{focus}|{share_year}": share_vals[(focus, s)]}
        for ps in psets:
            key = f"share|{ps.code}|{share_year}"
            r[key] = peer_cells.get((key, s), float("nan"))
        table_rows.append(r)
    data = _assemble(table_rows, columns, ["section", "row"])

    caveats: list[Caveat] = _modelled_caveat([res], [panel])
    if missing:
        caveats.append(Caveat("table_cells_missing", "info", f"{len(missing)} focus-country cells could not be computed.",
                              {"cells": _limited(missing), "n": len(missing)}))
    if focus_in_peers:
        caveats.append(Caveat("focus_in_peer_set", "warn", f"The focus country {focus} is a member of peer set(s) "
                              f"{focus_in_peers}; its own values enter the peer statistic.", {"peer_sets": focus_in_peers}))
    uncovered = _uncovered_changes(windows)
    if uncovered:
        caveats.append(Caveat("window_change_not_covered", "info",
                              "Growth windows are measured between end-point levels, so these changes fall in no "
                              f"column: {', '.join(_limited(uncovered))} (e.g. '1991-2009' ends at the 2009 level and "
                              "'2010-2022' starts at the 2010 level).",
                              {"changes": _limited(uncovered), "n": len(uncovered),
                               "windows": {k: list(v) for k, v in windows.items()}}))
    if stat == "median":
        caveats.append(Caveat("peer_median_shares", "info", "Peer medians of sector shares need not sum to 100.", {}))
    caveats += _endpoint_caveats(list(windows.values()) + [(share_year, share_year)], statuses,
                                 "standardized_employment_table")
    notes = [f"Annualized growth: {'CAGR' if growth == 'cagr' else 'log growth'} between window end-point levels "
             "(a window start-end covers the growth years after start up to end); a single year is growth over the "
             "previous year.",
             f"Peer columns: {stat} across members with data."]
    if source_note:
        notes.insert(0, source_note)
    layout = {
        "kind": "employment",
        "title": title or f"Standardized employment table, {focus_name}",
        "label_columns": ["row"],
        "label_headers": {"row": "Sector"},
        "sections": [{"key": "growth", "title": "Annualized employment growth (%)"},
                     {"key": "share", "title": f"Share of total employment, {share_year} (%)"}],
        "columns": columns,
        "decimals": 1,
        "notes": notes,
    }
    params = {"focus": focus, "peer_sets": [ps.to_dict() for ps in psets],
              "periods": {k: list(v) for k, v in windows.items()}, "peer_periods": plabels,
              "share_year": share_year, "sectors": dict(sectors) if sectors else None, "total": total,
              "stat": stat, "growth": growth, "min_n": min_n}
    out = _done([res, *peer_results], data, "standardized_employment_table", params, caveats, notes,
                module="standard_tables")
    out.meta["table"] = layout
    out.meta["cells"] = _cells(data, columns, peer_n)
    return out


# ---------------------------------------------------------------------------
# macro table
# ---------------------------------------------------------------------------
@dataclass
class MacroRow:
    """One row of the standardized macro table.

    ``agg`` per window ``(start, end)``: ``"mean"`` (average of the
    observations dated ``start..end``), ``"cagr"`` (compound annual growth
    over the growth years ``start..end``, i.e. between the levels of
    ``start - 1`` and ``end``; a single period means growth over the previous
    period), ``"first"`` / ``"last"`` (value at the window's start / end) or
    ``"sum"``. ``indicator`` filters a multi-indicator frame. ``source`` is
    the text of the table's Source column (default: the frame's dataset).
    """

    label: str
    frame: Frame
    agg: Literal["mean", "cagr", "first", "last", "sum"] = "mean"
    section: str = ""
    indicator: str | None = None
    source: str = ""
    decimals: int | None = None


def _cagr_levels(lo: str, hi: str) -> tuple[str, str]:
    """Level periods for a CAGR over the inclusive growth years ``lo..hi``: ``(lo - 1 period, hi)``."""
    plo, phi = _window((lo, hi))
    return format_period(plo.freq, plo.ordinal - 1), phi.text


def _window_value(series: pd.DataFrame, iso3: str, lo: str, hi: str, agg: str
                  ) -> tuple[float, list[str], int, int, str]:
    """Aggregate one country's series over a window.

    Returns ``(value, statuses, n_used, n_expected, reason)``; ``reason`` says
    why ``value`` is NaN (empty otherwise).
    """
    if iso3 not in series.index.get_level_values(0):
        return float("nan"), [], 0, 0, "no data for this country in the row"
    s = series.loc[iso3]

    def has(p: str) -> bool:
        return p in s.index and bool(np.isfinite(float(s.loc[p, "value"])))

    if agg == "cagr":
        a, b = _cagr_levels(lo, hi)
        absent = [p for p in (a, b) if not has(p)]
        if absent:
            return float("nan"), [], 0, 2, (f"no level for {absent} (the CAGR over growth years {lo}-{hi} uses the "
                                            f"levels of {a} and {b})")
        x0, xT = float(s.loc[a, "value"]), float(s.loc[b, "value"])
        v = _growth(x0, xT, _years_between(a, b), "cagr")
        ok = bool(np.isfinite(v))
        return (v, [str(s.loc[a, "obs_status"]), str(s.loc[b, "obs_status"])], 2 if ok else 0, 2,
                "" if ok else f"non-positive level in {a} or {b}")
    if agg in ("first", "last"):
        p = lo if agg == "first" else hi
        if not has(p):
            return float("nan"), [], 0, 1, f"no value in {p}"
        return float(s.loc[p, "value"]), [str(s.loc[p, "obs_status"])], 1, 1, ""
    plo, phi = _window(lo if lo == hi else (lo, hi))
    inside = []
    for p, row in s.iterrows():
        try:
            pp = parse_period(p)
        except ValueError:
            continue
        if pp.freq == plo.freq and plo.ordinal <= pp.ordinal <= phi.ordinal and np.isfinite(row["value"]):
            inside.append((float(row["value"]), str(row["obs_status"])))
    expected = phi.ordinal - plo.ordinal + 1
    if not inside:
        return float("nan"), [], 0, expected, f"no observation in {lo}-{hi}" if lo != hi else f"no value in {lo}"
    vals = [v for v, _ in inside]
    v = float(np.mean(vals)) if agg == "mean" else float(np.sum(vals))
    return v, [x for _, x in inside], len(inside), expected, ""


def _macro_windows(periods: Mapping[str, PeriodSpec]) -> dict[str, tuple[str, str]]:
    out = {}
    for label, spec in periods.items():
        if isinstance(spec, (tuple, list)):
            lo, hi = _window(spec)
            out[str(label)] = (lo.text, hi.text)
        else:
            p = parse_period(str(spec).strip())
            out[str(label)] = (p.text, p.text)
    return out


def standardized_macro_table(
    frames: Mapping[str, Frame | MacroRow] | Sequence[MacroRow],
    focus: str,
    peers: PeerSets,
    periods: Mapping[str, PeriodSpec],
    *,
    peer_periods: Sequence[str] | None = None,
    stat: Stat = "median",
    min_n: int = 3,
    require_complete: bool = False,
    title: str | None = None,
    show_source: bool = True,
    source_note: str = "",
) -> Result:
    """Standardized macro table: period averages for the focus country vs peer statistics.

    ``frames`` maps row labels to frames (period average of a single
    indicator, e.g. real GDP growth from WDI) or to :class:`MacroRow` specs
    (aggregation, section, source text, decimals), in row order.
    ``periods``: ``{label: (start, end) | period}`` inclusive windows, e.g.
    ``{"2000-2013": ("2000", "2013"), "2014-2024": ("2014", "2024")}``;
    ``peer_periods`` (labels, default all) are shown for each peer set.

    A focus cell averages the observations inside the window (or applies the
    row's ``agg``; a ``"cagr"`` row compounds the growth of the same years,
    from the level of ``start - 1`` to the level of ``end``, see the module
    docstring); with ``require_complete`` windows with gaps are left empty.
    Every focus cell that stays empty is listed, with the reason, in a
    ``table_cells_missing`` info caveat. Peer cells are the ``stat`` across
    members of the members' own window values. Forecast observations inside a
    window raise a ``forecast_in_window`` caveat. A row whose periods (focus
    and peer members) mix frequencies, or differ from the windows' frequency
    (fiscal-year data with calendar-year windows), raises
    :class:`~cgjrkit.methods.harmonize.FrequencyMismatchError`.
    """
    _check_stat(stat)
    if not periods:
        raise ValueError("periods must name at least one window")
    focus = str(focus).strip().upper()
    psets = _peersets(peers)
    windows = _macro_windows(periods)
    plabels = list(windows) if peer_periods is None else [str(p) for p in peer_periods]
    unknown = [p for p in plabels if p not in windows]
    if unknown:
        raise ValueError(f"peer_periods {unknown} are not labels of periods {list(windows)}")
    if isinstance(frames, Mapping):
        specs = [f if isinstance(f, MacroRow) else MacroRow(label=str(lab), frame=f) for lab, f in frames.items()]
        for lab, spec in zip(frames, specs):
            if isinstance(frames[lab], MacroRow) and not spec.label:
                spec.label = str(lab)
    else:
        specs = list(frames)
    if not specs:
        raise ValueError("frames must define at least one row")
    labels = [s.label for s in specs]
    if len(set(labels)) != len(labels):
        raise ValueError(f"row labels must be unique, got {labels}")
    window_freqs = {lab: parse_period(lo).freq for lab, (lo, _) in windows.items()}
    relevant = {focus} | {m for ps in psets for m in ps.members}

    inputs: list[Result] = []
    columns: list[dict] = []
    focus_name = focus
    table_rows: list[dict] = []
    incomplete: list[dict] = []
    missing_cells: list[dict] = []
    forecast: list[str] = []
    statuses: list[str] = []
    peer_rows: dict[str, list[dict]] = {ps.code: [] for ps in psets}
    row_decimals: dict[str, int] = {}
    for spec in specs:
        if spec.agg not in ("mean", "cagr", "first", "last", "sum"):
            raise ValueError(f"row {spec.label!r}: unknown agg {spec.agg!r}")
        r_in, series = _single_series(spec.frame, f"row {spec.label!r}", spec.indicator)
        inputs.append(r_in)
        used_periods = series.index[series.index.get_level_values(0).isin(relevant)].get_level_values(1)
        _check_frequency(list(used_periods), window_freqs, f"row {spec.label!r}")
        if focus_name == focus:
            focus_name = _country_name(r_in.data, focus)
        source = spec.source or next((p.dataset or p.source for p in r_in.provenance if p.source != "cgjrkit"), "") \
            or ", ".join(sorted({d for d in r_in.data["dataset"] if d})) or ", ".join(sorted({d for d in r_in.data["source"] if d}))
        if spec.decimals is not None:
            row_decimals[spec.label] = spec.decimals
        row: dict[str, Any] = {"section": spec.section, "row": spec.label, "source": source}
        for label, (lo, hi) in windows.items():
            v, st, n, expected, reason = _window_value(series, focus, lo, hi, spec.agg)
            statuses += st
            if "forecast" in st:
                forecast.append(f"{spec.label}@{label}")
            if spec.agg in ("mean", "sum") and 0 < n < expected:
                incomplete.append({"row": spec.label, "window": label, "n": n, "expected": expected})
                if require_complete:
                    v = float("nan")
                    reason = f"{n} of {expected} observations (require_complete=True)"
            if not np.isfinite(v):
                missing_cells.append({"row": spec.label, "window": label, "agg": spec.agg, "reason": reason})
            row[f"{focus}|{label}"] = v
        for ps in psets:
            for m in ps.members:
                for label in plabels:
                    lo, hi = windows[label]
                    v, st, n, expected, _ = _window_value(series, m, lo, hi, spec.agg)
                    if require_complete and spec.agg in ("mean", "sum") and n < expected:
                        continue
                    if np.isfinite(v):
                        peer_rows[ps.code].append(_member_row(m, spec.label, label, v, "", combine_status(st)))
                        if "forecast" in st:
                            forecast.append(f"{spec.label}@{label} ({m})")
        table_rows.append(row)

    for label in windows:
        columns.append({"key": f"{focus}|{label}", "group": focus_name, "sub": label, "sections": None,
                        "entity": focus, "kind": "focus"})
    peer_results: list[Result] = []
    peer_n: dict[tuple[str, str], int] = {}
    for ps in psets:
        for label in plabels:
            columns.append({"key": f"{ps.code}|{label}", "group": f"{ps.name} ({stat})", "sub": label,
                            "sections": None, "entity": ps.code, "kind": "peer"})
        vals, ns, agg = _peer_values(peer_rows[ps.code], ps, stat, min_n)
        if agg is None:
            agg = Result(data=pd.DataFrame(), caveats=[Caveat(
                "peer_no_data", "warn", f"No member of peer set {ps.name!r} has data for these rows and windows.",
                {"peer_set": ps.code, "members": ps.members})])
        peer_results.append(agg)
        for row in table_rows:
            for label in plabels:
                key = f"{ps.code}|{label}"
                row[key] = vals.get((row["row"], label), float("nan"))
                if (row["row"], label) in ns:
                    peer_n[(key, row["row"])] = ns[(row["row"], label)]

    label_cols = ["section", "row", "source"]
    data = _assemble(table_rows, columns, label_cols)
    caveats: list[Caveat] = _modelled_caveat(inputs, [r.data for r in inputs])
    if incomplete:
        caveats.append(Caveat("period_average_incomplete", "warn",
                              f"{len(incomplete)} focus-country window values use fewer observations than the window length"
                              + (" and were left empty." if require_complete else "."),
                              {"cells": _limited(incomplete), "n": len(incomplete)}))
    if missing_cells:
        caveats.append(Caveat("table_cells_missing", "info",
                              f"{len(missing_cells)} focus-country cells could not be computed.",
                              {"cells": _limited(missing_cells), "n": len(missing_cells)}))
    if forecast:
        caveats.append(Caveat("forecast_in_window", "warn", "Some windows include forecast values.",
                              {"at": _limited(forecast), "n": len(forecast)}))
    caveats += _endpoint_caveats(list(windows.values()), [s for s in statuses if s != "forecast"],
                                 "standardized_macro_table")
    notes = ["Columns are period averages (or the row's stated aggregation) over inclusive windows."]
    if any(spec.agg == "cagr" for spec in specs):
        notes.append("CAGR rows compound the growth of the same years: from the level of the year before the window "
                     "starts to the level of its last year.")
    notes.append(f"Peer columns: {stat} across members of the members' window values.")
    if source_note:
        notes.insert(0, source_note)
    sections = []
    for s in [r["section"] for r in table_rows]:
        if s not in [x["key"] for x in sections]:
            sections.append({"key": s, "title": s})
    layout = {
        "kind": "macro",
        "title": title or f"Standardized macro table, {focus_name}",
        "label_columns": ["row", "source"] if show_source else ["row"],
        "label_headers": {"row": "Indicator", "source": "Source"},
        "sections": sections,
        "columns": columns,
        "decimals": 1,
        "row_decimals": row_decimals,
        "notes": notes,
    }
    params = {"rows": [{"label": s.label, "agg": s.agg, "section": s.section, "indicator": s.indicator}
                       for s in specs],
              "focus": focus, "peer_sets": [ps.to_dict() for ps in psets],
              "periods": {k: list(v) for k, v in windows.items()}, "peer_periods": plabels, "stat": stat,
              "min_n": min_n, "require_complete": require_complete}
    out = _done([*inputs, *peer_results], data, "standardized_macro_table", params, caveats, notes,
                module="standard_tables")
    out.meta["table"] = layout
    out.meta["cells"] = _cells(data, columns, peer_n)
    return out


# ---------------------------------------------------------------------------
# export
# ---------------------------------------------------------------------------
def _layout(table: Result) -> dict:
    layout = table.meta.get("table")
    if not isinstance(layout, Mapping) or "columns" not in layout:
        raise ValueError("table has no layout; build it with standardized_employment_table or standardized_macro_table")
    return dict(layout)


def _blocks(table: Result, layout: Mapping) -> tuple[bool, list[tuple[dict, pd.DataFrame, list[dict]]]]:
    """``(shared_header, [(section, rows, columns), ...])`` in display order."""
    data = table.data
    columns = list(layout["columns"])
    shared = all(c.get("sections") is None for c in columns)
    blocks = []
    for sec in layout["sections"]:
        rows = data[data["section"] == sec["key"]]
        if rows.empty:
            continue
        cols = columns if shared else [c for c in columns if c.get("sections") is None or sec["key"] in c["sections"]]
        blocks.append((sec, rows, cols))
    return shared, blocks


def _groups(cols: list[dict]) -> list[tuple[str, int]]:
    """Consecutive column groups -> ``[(group label, span), ...]``."""
    out: list[tuple[str, int]] = []
    for c in cols:
        if out and out[-1][0] == c["group"]:
            out[-1] = (out[-1][0], out[-1][1] + 1)
        else:
            out.append((c["group"], 1))
    return out


def _source_line(table: Result) -> str:
    """Source line from the data providers' provenance (cgjrkit's own entries are not sources)."""
    cites: list[str] = []
    for p in table.provenance:
        if p.source == "cgjrkit":
            continue
        cite = p.citation or p.dataset or p.source
        if p.vintage:
            cite = f"{cite} ({p.vintage})"
        if cite and cite not in cites:
            cites.append(cite)
    return ("Source: " + "; ".join(cites) + ".") if cites else ""


def _footer_lines(table: Result, layout: Mapping) -> list[str]:
    lines = []
    src = _source_line(table)
    if src:
        lines.append(src)
    lines += [f"Note: {n}" for n in layout.get("notes", [])]
    lines += [f"Caveat ({c.severity}): {c.message}" for c in table.caveats if c.severity in ("warn", "block")]
    return lines


def _fmt(v: Any, decimals: int) -> str:
    if v is None or (isinstance(v, float) and not np.isfinite(v)):
        return _MISSING
    return f"{float(v):.{decimals}f}"


def to_xlsx(table: Result, path: str | Path, *, sheet_name: str = "Table") -> Path:
    """Write a standardized table to an XLSX file with formatted headers (openpyxl).

    Header rows are bold on a light fill with column groups merged across
    their periods; section rows are bold on a lighter fill; numbers keep full
    precision with a ``0.0`` display format (per-row decimals in the macro
    table); missing cells show ``".."``. The source line, notes and WARN /
    BLOCK caveats are written below the table.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    layout = _layout(table)
    shared, blocks = _blocks(table, layout)
    label_cols = list(layout["label_columns"])
    headers = layout.get("label_headers", {})
    decimals = int(layout.get("decimals", 1))
    row_dec = dict(layout.get("row_decimals", {}))
    nl = len(label_cols)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor=_HEADER_FILL)
    sec_fill = PatternFill("solid", fgColor=_SECTION_FILL)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="888888")

    ws.cell(row=1, column=1, value=layout.get("title", "")).font = Font(bold=True, size=13)
    r = 3

    def header(r0: int, cols: list[dict]) -> int:
        for j, lc in enumerate(label_cols, start=1):
            ws.merge_cells(start_row=r0, start_column=j, end_row=r0 + 1, end_column=j)
            cell = ws.cell(row=r0, column=j, value=headers.get(lc, lc))
            cell.font, cell.fill, cell.alignment = bold, head_fill, Alignment(vertical="center", wrap_text=True)
            ws.cell(row=r0 + 1, column=j).fill = head_fill
        col = nl + 1
        for group, span in _groups(cols):
            if span > 1:
                ws.merge_cells(start_row=r0, start_column=col, end_row=r0, end_column=col + span - 1)
            cell = ws.cell(row=r0, column=col, value=group)
            cell.font, cell.fill, cell.alignment = bold, head_fill, center
            for k in range(span):
                ws.cell(row=r0, column=col + k).fill = head_fill
            col += span
        for k, c in enumerate(cols):
            cell = ws.cell(row=r0 + 1, column=nl + 1 + k, value=c["sub"])
            cell.font, cell.fill, cell.alignment = bold, head_fill, center
        for j in range(1, nl + len(cols) + 1):
            ws.cell(row=r0 + 1, column=j).border = Border(bottom=thin)
        return r0 + 2

    max_cols = nl + max((len(cols) for _, _, cols in blocks), default=0)
    first_data_row = None
    if shared and blocks:
        r = header(r, blocks[0][2])
        first_data_row = r
    for sec, rows, cols in blocks:
        if not shared:
            r = header(r, cols)
        if sec.get("title"):
            cell = ws.cell(row=r, column=1, value=sec["title"])
            cell.font = Font(bold=True, italic=True)
            for j in range(1, nl + len(cols) + 1):
                ws.cell(row=r, column=j).fill = sec_fill
            r += 1
        for _, row in rows.iterrows():
            for j, lc in enumerate(label_cols, start=1):
                ws.cell(row=r, column=j, value=row[lc])
            d = row_dec.get(row["row"], decimals)
            for k, c in enumerate(cols):
                v = row[c["key"]]
                cell = ws.cell(row=r, column=nl + 1 + k)
                if pd.isna(v):
                    cell.value = _MISSING
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.value = float(v)
                    cell.number_format = "0" if d == 0 else "0." + "0" * d
            r += 1
        if not shared:
            r += 1
    r += 1
    for line in _footer_lines(table, layout):
        ws.cell(row=r, column=1, value=line).font = Font(italic=True, size=9)
        r += 1
    ws.column_dimensions["A"].width = 44
    for j in range(2, max_cols + 1):
        ws.column_dimensions[get_column_letter(j)].width = 18 if j <= nl else 13
    if first_data_row:
        ws.freeze_panes = ws.cell(row=first_data_row, column=nl + 1)
    p = Path(path)
    wb.save(p)
    return p


_CSS = """
.cgjr-table-wrap { font-family: Helvetica, Arial, sans-serif; font-size: 13px; color: #1f2933; max-width: 100%; overflow-x: auto; }
.cgjr-table { border-collapse: collapse; margin: 0 0 1rem 0; }
.cgjr-table caption { text-align: left; font-weight: bold; font-size: 15px; padding: 0 0 6px 0; }
.cgjr-table th, .cgjr-table td { padding: 4px 8px; border-bottom: 1px solid #e5e8eb; }
.cgjr-table thead th { background: #d9e2ea; text-align: center; vertical-align: middle; }
.cgjr-table thead th.label { text-align: left; }
.cgjr-table td.num { text-align: right; font-variant-numeric: tabular-nums; white-space: nowrap; }
.cgjr-table tr.section th { background: #eef1f4; text-align: left; font-style: italic; }
.cgjr-table tbody th { text-align: left; font-weight: normal; }
.cgjr-notes { font-size: 11px; color: #555; margin: 0.25rem 0; }
""".strip()


def to_html(table: Result, path: str | Path | None = None, *, full_document: bool = False,
            include_style: bool = True) -> str:
    """Render a standardized table as HTML (and write it to ``path`` if given).

    Returns a ``<div class="cgjr-table-wrap">`` fragment (with a small
    scoped ``<style>`` block unless ``include_style=False``), or a complete
    UTF-8 document with ``full_document=True``. Text is HTML-escaped; missing
    cells show ``".."``.
    """
    layout = _layout(table)
    shared, blocks = _blocks(table, layout)
    label_cols = list(layout["label_columns"])
    headers = layout.get("label_headers", {})
    decimals = int(layout.get("decimals", 1))
    row_dec = dict(layout.get("row_decimals", {}))
    esc = html.escape
    title = str(layout.get("title", ""))

    def thead(cols: list[dict]) -> list[str]:
        out = ["<thead>", "<tr>"]
        out += [f'<th class="label" rowspan="2" scope="col">{esc(str(headers.get(lc, lc)))}</th>' for lc in label_cols]
        out += [f'<th colspan="{span}" scope="colgroup">{esc(str(g))}</th>' for g, span in _groups(cols)]
        out += ["</tr>", "<tr>"]
        out += [f'<th scope="col">{esc(str(c["sub"]))}</th>' for c in cols]
        out += ["</tr>", "</thead>"]
        return out

    def body(sec: dict, rows: pd.DataFrame, cols: list[dict], with_title: bool) -> list[str]:
        out = ["<tbody>"]
        if with_title and sec.get("title"):
            out.append(f'<tr class="section"><th colspan="{len(label_cols) + len(cols)}" scope="colgroup">'
                       f'{esc(str(sec["title"]))}</th></tr>')
        for _, row in rows.iterrows():
            d = row_dec.get(row["row"], decimals)
            cells = [f'<th scope="row">{esc(str(row[label_cols[0]]))}</th>']
            cells += [f"<td>{esc(str(row[lc]))}</td>" for lc in label_cols[1:]]
            cells += [f'<td class="num">{esc(_fmt(row[c["key"]], d))}</td>' for c in cols]
            out.append("<tr>" + "".join(cells) + "</tr>")
        out.append("</tbody>")
        return out

    parts = ['<div class="cgjr-table-wrap">']
    if include_style or full_document:
        parts.append(f"<style>\n{_CSS}\n</style>")
    if shared:
        parts.append('<table class="cgjr-table">')
        parts.append(f"<caption>{esc(title)}</caption>")
        if blocks:
            parts += thead(blocks[0][2])
        for sec, rows, cols in blocks:
            parts += body(sec, rows, cols, with_title=True)
        parts.append("</table>")
    else:
        parts.append(f'<p class="cgjr-title"><strong>{esc(title)}</strong></p>')
        for sec, rows, cols in blocks:
            parts.append('<table class="cgjr-table">')
            parts.append(f"<caption>{esc(str(sec.get('title', '')))}</caption>")
            parts += thead(cols)
            parts += body(sec, rows, cols, with_title=False)
            parts.append("</table>")
    for line in _footer_lines(table, layout):
        parts.append(f'<p class="cgjr-notes">{esc(line)}</p>')
    parts.append("</div>")
    fragment = "\n".join(parts)
    if full_document:
        fragment = ("<!doctype html>\n<html lang=\"en\">\n<head>\n<meta charset=\"utf-8\">\n"
                    "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
                    f"<title>{esc(title)}</title>\n</head>\n<body>\n{fragment}\n</body>\n</html>\n")
    if path is not None:
        Path(path).write_text(fragment, encoding="utf-8")
    return fragment
