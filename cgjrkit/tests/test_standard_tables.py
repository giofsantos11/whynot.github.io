"""Unit tests for cgjrkit.methods.standard_tables (synthetic values, hand-computed cells).

Employment levels (thousands, fake ISO3 codes; AAA is the focus country):

======  ======================  ======================
        Agriculture             Services
        2000   2002   2003      2000   2002   2003
======  ======================  ======================
AAA      100    121  133.1       100    100    100
BBB      100    100    100       100    144    144
CCC      200    200    220       100    121    121
DDD       50     72     72       150    150    165
======  ======================  ======================

Windows: "2000-02" = CAGR 2000->2002; "2003" = growth 2002->2003.
AAA agriculture: 10% and 10%; peers (BBB, CCC, DDD) agriculture 2000-02: 0, 0, 20
-> median 0, mean 20/3; services 2000-02: 20, 10, 0 -> median 10.
"""

from __future__ import annotations

import math

import pandas as pd
import pytest
from openpyxl import load_workbook

from cgjrkit.core.provenance import Provenance, Result
from cgjrkit.methods.peers import PeerSet
from cgjrkit.methods.structural_change import StructuralChangeError
from cgjrkit.methods.standard_tables import (
    MacroRow,
    standardized_employment_table,
    standardized_macro_table,
    to_html,
    to_xlsx,
)

LEVELS = {
    "AAA": {"Agr": {2000: 100, 2002: 121, 2003: 133.1}, "Srv": {2000: 100, 2002: 100, 2003: 100}},
    "BBB": {"Agr": {2000: 100, 2002: 100, 2003: 100}, "Srv": {2000: 100, 2002: 144, 2003: 144}},
    "CCC": {"Agr": {2000: 200, 2002: 200, 2003: 220}, "Srv": {2000: 100, 2002: 121, 2003: 121}},
    "DDD": {"Agr": {2000: 50, 2002: 72, 2003: 72}, "Srv": {2000: 150, 2002: 150, 2003: 165}},
}
SECTORS = {"EMP_Agr": "Agriculture", "EMP_Srv": "Services"}
PERIODS = {"2000-02": ("2000", "2002"), "2003": "2003"}
STRUCT = PeerSet(name="Structural peers", code="STRUCT", members=["BBB", "CCC", "DDD"],
                 criteria="synthetic", approver="Test approver", frozen_at="2026-03-01")


def emp_frame(status: str = "") -> pd.DataFrame:
    rows = []
    for iso, d in LEVELS.items():
        for s, ys in d.items():
            for y, v in ys.items():
                rows.append({"source": "synthetic", "dataset": "TEST", "indicator": f"EMP_{s}", "indicator_label": "",
                             "iso3": iso, "country": "Aland" if iso == "AAA" else iso, "period": str(y),
                             "value": v, "unit": "thousands", "obs_status": status})
    return pd.DataFrame(rows)


def cell(t: Result, section: str, row: str, key: str) -> float:
    hit = t.data[(t.data["section"] == section) & (t.data["row"] == row)]
    assert len(hit) == 1
    return float(hit[key].iloc[0])


# ---------------------------------------------------------------------------
# employment table
# ---------------------------------------------------------------------------
def test_employment_table_growth_and_shares_by_hand():
    t = standardized_employment_table(emp_frame(), "AAA", [STRUCT], PERIODS, sectors=SECTORS)
    assert list(t.data.columns) == ["section", "row", "growth|AAA|2000-02", "growth|AAA|2003",
                                    "growth|STRUCT|2000-02", "growth|STRUCT|2003", "share|AAA|2003",
                                    "share|STRUCT|2003"]
    assert list(t.data["row"]) == ["Agriculture", "Services", "Total", "Agriculture", "Services"]
    assert cell(t, "growth", "Agriculture", "growth|AAA|2000-02") == pytest.approx(10.0, rel=1e-12)
    assert cell(t, "growth", "Agriculture", "growth|AAA|2003") == pytest.approx(10.0, rel=1e-12)
    assert cell(t, "growth", "Total", "growth|AAA|2000-02") == pytest.approx(100 * (math.sqrt(221 / 200) - 1))
    assert cell(t, "growth", "Total", "growth|AAA|2003") == pytest.approx(100 * (233.1 / 221 - 1))
    assert cell(t, "growth", "Agriculture", "growth|STRUCT|2000-02") == 0.0
    assert cell(t, "growth", "Services", "growth|STRUCT|2000-02") == pytest.approx(10.0)
    # peer totals 2000-02: BBB 200->244, CCC 300->321, DDD 200->222 -> median = DDD
    assert cell(t, "growth", "Total", "growth|STRUCT|2000-02") == pytest.approx(100 * (math.sqrt(222 / 200) - 1))
    # shares in 2003 (default share year = end of the last window)
    assert cell(t, "share", "Agriculture", "share|AAA|2003") == pytest.approx(100 * 133.1 / 233.1)
    assert cell(t, "share", "Agriculture", "share|STRUCT|2003") == pytest.approx(100 * 100 / 244)  # median of 41.0, 64.5, 30.4
    assert math.isnan(cell(t, "share", "Agriculture", "growth|AAA|2003"))
    n = {(c["column"], c["row"]): c.get("n") for c in t.meta["cells"]}
    assert n[("growth|STRUCT|2000-02", "Agriculture")] == 3 and n[("growth|AAA|2003", "Agriculture")] is None
    layout = t.meta["table"]
    assert layout["kind"] == "employment" and [s["key"] for s in layout["sections"]] == ["growth", "share"]
    assert t.provenance[-1].tool == "standard_tables.standardized_employment_table"
    assert any(p.tool == "peers.peer_aggregate" for p in t.provenance)
    assert {c.code for c in t.caveats} == {"peer_median_shares"}


def test_employment_table_mean_peer_periods_and_caveats():
    asp = {"Aspirational": ["BBB", "AAA"]}  # unfrozen, includes the focus, only 2 members
    t = standardized_employment_table(emp_frame(status="modelled"), "AAA", {**{"Structural": STRUCT}, **asp},
                                      PERIODS, sectors=SECTORS, stat="mean", peer_periods=["2000-02"])
    assert "growth|STRUCT|2003" not in t.data.columns
    assert cell(t, "growth", "Agriculture", "growth|STRUCT|2000-02") == pytest.approx(20 / 3)
    # mean shares sum to 100 across sectors
    shares = t.data[t.data["section"] == "share"]["share|STRUCT|2003"].sum()
    assert shares == pytest.approx(100.0)
    codes = {c.code for c in t.caveats}
    assert {"peer_set_not_frozen", "peer_small_n", "focus_in_peer_set", "modelled_inputs"} <= codes
    assert "peer_median_shares" not in codes


def test_employment_table_total_sector_share_year_and_errors():
    df = emp_frame()
    tot = df.groupby(["iso3", "period", "country"], as_index=False)["value"].sum()
    tot = tot.assign(source="synthetic", dataset="TEST", indicator="EMP_T", indicator_label="", unit="thousands",
                     obs_status="")
    full = pd.concat([df, tot], ignore_index=True)
    t = standardized_employment_table(full, "AAA", [STRUCT], PERIODS, sectors={**SECTORS, "EMP_T": "Total"},
                                      total="Total", share_year="2002")
    assert "share|AAA|2002" in t.data.columns
    assert cell(t, "share", "Agriculture", "share|AAA|2002") == pytest.approx(100 * 121 / 221)
    assert cell(t, "growth", "Total", "growth|AAA|2000-02") == pytest.approx(100 * (math.sqrt(221 / 200) - 1))
    with pytest.raises(ValueError):
        standardized_employment_table(df, "AAA", [STRUCT], PERIODS, sectors=SECTORS, peer_periods=["1990s"])
    with pytest.raises(StructuralChangeError):
        standardized_employment_table(df, "ZZZ", [STRUCT], PERIODS, sectors=SECTORS)
    with pytest.raises(ValueError):
        standardized_employment_table(df, "AAA", [STRUCT], PERIODS, sectors=SECTORS, stat="p75")


# ---------------------------------------------------------------------------
# macro table
# ---------------------------------------------------------------------------
def macro_frames():
    growth = {"AAA": [1, 2, 3, 4], "BBB": [2, 2, 2, 2], "CCC": [4, 4, 6, 6], "DDD": [0, 0, 1, 1]}
    g = pd.DataFrame([{"source": "worldbank", "dataset": "WDI", "indicator": "NY.GDP.MKTP.KD.ZG", "iso3": i,
                       "country": "Aland" if i == "AAA" else i, "period": str(y), "value": v, "unit": "%",
                       "indicator_label": "GDP growth", "obs_status": "forecast" if (i == "AAA" and y == 2003) else ""}
                      for i, vals in growth.items() for y, v in zip(range(2000, 2004), vals)])
    levels = {"AAA": [100, 110, 121, 133.1], "BBB": [100, 100, 100, 100], "CCC": [100, 120, 120, 144],
              "DDD": [100, 105, 105, 105]}
    lv = pd.DataFrame([{"source": "worldbank", "dataset": "WDI", "indicator": "NY.GDP.PCAP.KD", "iso3": i,
                        "country": i, "period": str(y), "value": v, "unit": "constant 2015 US$"}
                       for i, vals in levels.items() for y, v in zip(range(2000, 2004), vals)])
    g_res = Result(data=g, provenance=[Provenance(source="worldbank", dataset="WDI", vintage="WDI 2026-07")])
    return g_res, lv


def test_macro_table_period_averages_cagr_and_peer_median():
    g, lv = macro_frames()
    rows = {
        "Real GDP growth (%)": MacroRow("Real GDP growth (%)", g, section="Growth"),
        "GDP per capita growth (%)": MacroRow("GDP per capita growth (%)", lv, agg="cagr", section="Growth",
                                              source="WDI (calc.)", decimals=2),
        "GDP per capita, end of period": MacroRow("GDP per capita, end of period", lv, agg="last",
                                                  section="Levels", decimals=0),
    }
    periods = {"2000-01": ("2000", "2001"), "2002-03": ("2002", "2003")}
    t = standardized_macro_table(rows, "AAA", [STRUCT], periods)
    assert list(t.data.columns) == ["section", "row", "source", "AAA|2000-01", "AAA|2002-03",
                                    "STRUCT|2000-01", "STRUCT|2002-03"]
    r = t.data.set_index("row")
    assert r.loc["Real GDP growth (%)", "AAA|2000-01"] == 1.5
    assert r.loc["Real GDP growth (%)", "AAA|2002-03"] == 3.5
    assert r.loc["Real GDP growth (%)", "STRUCT|2000-01"] == 2.0      # median of 2, 4, 0
    assert r.loc["Real GDP growth (%)", "STRUCT|2002-03"] == 2.0      # median of 2, 6, 1
    assert r.loc["Real GDP growth (%)", "source"] == "WDI"
    # CAGR over the growth years of the window = levels of (start - 1) and end:
    # "2002-03" -> 2001 to 2003: AAA 110->133.1 = 10%; BBB 0; CCC 120->144 = 100*(sqrt(1.2)-1) = 9.54; DDD 0
    assert r.loc["GDP per capita growth (%)", "AAA|2002-03"] == pytest.approx(10.0)
    assert r.loc["GDP per capita growth (%)", "STRUCT|2002-03"] == pytest.approx(0.0)   # median of 0, 9.54, 0
    # "2000-01" needs the 1999 level, which nobody has: empty, and reported
    assert math.isnan(r.loc["GDP per capita growth (%)", "AAA|2000-01"])
    assert math.isnan(r.loc["GDP per capita growth (%)", "STRUCT|2000-01"])
    missing = next(c for c in t.caveats if c.code == "table_cells_missing")
    assert missing.evidence["cells"] == [{"row": "GDP per capita growth (%)", "window": "2000-01", "agg": "cagr",
                                          "reason": "no level for ['1999'] (the CAGR over growth years 2000-2001 "
                                                    "uses the levels of 1999 and 2001)"}]
    assert r.loc["GDP per capita, end of period", "AAA|2002-03"] == pytest.approx(133.1)
    assert r.loc["GDP per capita, end of period", "STRUCT|2002-03"] == pytest.approx(105.0)
    assert "forecast_in_window" in {c.code for c in t.caveats}
    assert t.meta["table"]["row_decimals"] == {"GDP per capita growth (%)": 2, "GDP per capita, end of period": 0}
    assert [s["key"] for s in t.meta["table"]["sections"]] == ["Growth", "Levels"]


def test_macro_table_incomplete_windows_and_mean_stat():
    g, _ = macro_frames()
    data = g.data[~((g.data["iso3"] == "AAA") & (g.data["period"] == "2001"))]
    periods = {"2000-01": ("2000", "2001")}
    t = standardized_macro_table({"Real GDP growth (%)": data}, "AAA", {"Peers": ["BBB", "CCC", "DDD"]}, periods,
                                 stat="mean")
    assert t.data.loc[0, "AAA|2000-01"] == 1.0
    assert t.data.loc[0, "PEERS|2000-01"] == 2.0
    assert "period_average_incomplete" in {c.code for c in t.caveats}
    strict = standardized_macro_table({"Real GDP growth (%)": data}, "AAA", {"Peers": ["BBB", "CCC", "DDD"]},
                                      periods, require_complete=True)
    assert math.isnan(strict.data.loc[0, "AAA|2000-01"])
    with pytest.raises(ValueError):
        standardized_macro_table({"x": data}, "AAA", [STRUCT], periods, peer_periods=["1990s"])


# ---------------------------------------------------------------------------
# export
# ---------------------------------------------------------------------------
def test_to_xlsx_employment_table_layout(tmp_path):
    t = standardized_employment_table(emp_frame(), "AAA", [STRUCT], PERIODS, sectors=SECTORS,
                                      source_note="Synthetic data for tests.")
    path = to_xlsx(t, tmp_path / "employment.xlsx")
    ws = load_workbook(path).active
    assert ws["A1"].value == "Standardized employment table, Aland" and ws["A1"].font.bold
    assert ws["A3"].value == "Sector" and ws["A3"].font.bold and ws["A3"].fill.fgColor.rgb.endswith("D9E2EA")
    assert ws["B3"].value == "Aland" and ws["D3"].value == "Structural peers (median)"
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert {"A3:A4", "B3:C3", "D3:E3"} <= merged
    assert [ws.cell(row=4, column=j).value for j in range(2, 6)] == ["2000-02", "2003", "2000-02", "2003"]
    assert ws["A5"].value == "Annualized employment growth (%)" and ws["A5"].font.bold
    assert ws["A6"].value == "Agriculture" and ws["B6"].value == pytest.approx(10.0)
    assert ws["B6"].number_format == "0.0"
    texts = [c.value for c in ws["A"] if isinstance(c.value, str)]
    assert "Share of total employment, 2003 (%)" in texts
    assert any(s.startswith("Note: Synthetic data for tests.") for s in texts)


def test_to_xlsx_macro_missing_cells_and_decimals(tmp_path):
    g, lv = macro_frames()
    rows = [MacroRow("Real GDP growth (%)", g, section="Growth"),
            MacroRow("GDP per capita, end of period", lv, agg="last", section="Levels", decimals=0)]
    t = standardized_macro_table(rows, "AAA", [STRUCT], {"1990s": ("1990", "1999"), "2002-03": ("2002", "2003")})
    ws = load_workbook(to_xlsx(t, tmp_path / "macro.xlsx")).active
    assert ws["A3"].value == "Indicator" and ws["B3"].value == "Source"
    assert ws["C3"].value == "Aland" and ws["C4"].value == "1990s"
    values = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
    r_g = values["Real GDP growth (%)"]
    assert ws.cell(row=r_g, column=3).value == ".."                    # no data in the 1990s
    assert ws.cell(row=r_g, column=4).value == pytest.approx(3.5)
    r_l = values["GDP per capita, end of period"]
    assert ws.cell(row=r_l, column=4).number_format == "0"
    assert ws.freeze_panes == "C5"
    assert any(isinstance(v, str) and v.startswith("Source: WDI (WDI 2026-07)") for v in values)


def test_to_html_escapes_and_structure(tmp_path):
    g, _ = macro_frames()
    t = standardized_macro_table({"Exports & imports <% GDP>": g}, "AAA", [STRUCT],
                                 {"1990s": ("1990", "1999"), "2000-01": ("2000", "2001")}, title="Table <A1>")
    out = to_html(t, tmp_path / "t.html")
    assert (tmp_path / "t.html").read_text(encoding="utf-8") == out
    assert "<caption>Table &lt;A1&gt;</caption>" in out
    assert "Exports &amp; imports &lt;% GDP&gt;" in out
    assert '<th colspan="2" scope="colgroup">Aland</th>' in out
    assert '<td class="num">..</td>' in out and '<td class="num">1.5</td>' in out
    assert out.startswith('<div class="cgjr-table-wrap">')
    doc = to_html(t, full_document=True)
    assert doc.startswith("<!doctype html>") and "<title>Table &lt;A1&gt;</title>" in doc
    emp = standardized_employment_table(emp_frame(), "AAA", [STRUCT], PERIODS, sectors=SECTORS)
    h = to_html(emp, include_style=False)
    assert "<style>" not in h and h.count("<table") == 2
    with pytest.raises(ValueError):
        to_html(Result(data=pd.DataFrame()))


# ---------------------------------------------------------------------------
# review fixes: one window convention per table, frequencies, empty focus cells
# ---------------------------------------------------------------------------
def test_macro_table_single_period_and_window_cagr_follow_the_same_rule():
    _, lv = macro_frames()
    rows = {"cagr": MacroRow("cagr", lv, agg="cagr")}
    t = standardized_macro_table(rows, "AAA", [STRUCT], {"2003": "2003", "2003-03": ("2003", "2003"),
                                                          "2002-03": ("2002", "2003")})
    r = t.data.iloc[0]
    # single period 2003 = window (2003, 2003) = levels 2002 -> 2003: CCC 120 -> 144 = 20%
    assert r["AAA|2003"] == pytest.approx(10.0) and r["AAA|2003-03"] == pytest.approx(10.0)
    assert r["STRUCT|2003"] == pytest.approx(r["STRUCT|2003-03"])
    assert r["AAA|2002-03"] == pytest.approx(10.0)                  # 110 -> 133.1 over two years
    assert any("CAGR rows compound the growth of the same years" in n for n in t.meta["table"]["notes"])


def test_macro_table_refuses_mixed_frequency_rows_and_reports_empty_cells():
    from cgjrkit.methods.harmonize import FrequencyMismatchError

    g, _ = macro_frames()
    fy_focus = g.data.copy()
    fy_focus.loc[fy_focus["iso3"] == "AAA", "period"] = [f"FY{y}/{(y + 1) % 100:02d}" for y in range(2000, 2004)]
    with pytest.raises(FrequencyMismatchError, match="mixes period frequencies"):
        standardized_macro_table({"Growth": fy_focus}, "AAA", [STRUCT], {"2000-01": ("2000", "2001")})
    # a country outside the focus and the peer sets does not count
    other = pd.concat([g.data, g.data[g.data["iso3"] == "BBB"].assign(iso3="ZZZ", period=lambda d: "FY" + d["period"]
                                                                       + "/" + (d["period"].astype(int) + 1).astype(str).str[-2:])])
    assert standardized_macro_table({"Growth": other}, "AAA", [STRUCT], {"2000-01": ("2000", "2001")}).data.loc[
        0, "AAA|2000-01"] == 1.5
    # the focus country has no data at all for this row
    t = standardized_macro_table({"Growth": g.data[g.data["iso3"] != "AAA"]}, "AAA", [STRUCT],
                                 {"2000-01": ("2000", "2001")})
    cav = next(c for c in t.caveats if c.code == "table_cells_missing")
    assert cav.evidence["cells"] == [{"row": "Growth", "window": "2000-01", "agg": "mean",
                                      "reason": "no data for this country in the row"}]


def test_employment_table_reports_changes_between_windows_and_checks_frequency():
    from cgjrkit.methods.harmonize import FrequencyMismatchError

    t = standardized_employment_table(emp_frame(), "AAA", [STRUCT], {"2000-02": ("2000", "2002"), "2003": "2003"},
                                      sectors=SECTORS)
    assert "window_change_not_covered" not in {c.code for c in t.caveats}    # 2000->2002 then 2002->2003
    gap = standardized_employment_table(emp_frame(), "AAA", [STRUCT], {"2000-01": ("2000", "2001"),
                                                                       "2002-03": ("2002", "2003")}, sectors=SECTORS)
    cav = next(c for c in gap.caveats if c.code == "window_change_not_covered")
    assert cav.severity == "info" and cav.evidence["changes"] == ["2001->2002"]
    fy = emp_frame()
    fy["period"] = "FY" + fy["period"] + "/" + (fy["period"].astype(int) + 1).astype(str).str[-2:]
    with pytest.raises(FrequencyMismatchError):
        standardized_employment_table(fy, "AAA", [STRUCT], PERIODS, sectors=SECTORS)
